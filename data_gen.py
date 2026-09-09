from openpyxl import load_workbook
import json,re
p='/mnt/data/Юнитка ИП Сайдали Раджабалии(1).xlsx'
wb=load_workbook(p,data_only=True,read_only=True)
ws=wb['загрузка отчета по ценам']
main=wb['Юнит-экономика по товарам']
cost={}
for r in main.iter_rows(min_row=4,values_only=True):
    if r[2] and isinstance(r[9],(int,float)): cost[str(r[2]).strip()]=float(r[9])
products=[]
for r in ws.iter_rows(min_row=2,values_only=True):
    if r[0]!='Clara' or not r[3]: continue
    cat=r[1] or 'Без категории'; wb_art=str(int(r[2])) if isinstance(r[2],(int,float)) else str(r[2] or '')
    supplier=str(r[3]); barcode=str(int(r[4])) if isinstance(r[4],(int,float)) else str(r[4] or '')
    price_raw=r[8]
    price=0
    if isinstance(price_raw,(int,float)): price=float(price_raw)
    else:
      nums=re.findall(r'\d+(?:[.,]\d+)?',str(price_raw)); price=float(nums[-1].replace(',','.')) if nums else 0
    discount=float(r[10] or 0)
    products.append(dict(article=supplier,wbArticle=wb_art,name=cat,category=cat,barcode=barcode,cost=cost.get(supplier,0),price=price,discount=discount,colors=[],sizes=['48','50','52','54','56','58'],stockWB=int(r[5] or 0),stockSeller=int(r[6] or 0),variants=[]))
print('products',len(products))
open('/mnt/data/CLARA_V2/data.js','w',encoding='utf8').write('const INITIAL_PRODUCTS='+json.dumps(products,ensure_ascii=False,separators=(',',':'))+';\n')
